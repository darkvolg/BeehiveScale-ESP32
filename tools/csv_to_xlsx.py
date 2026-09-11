# -*- coding: utf-8 -*-
"""BeehiveScale: CSV из вкладки «Архив» → цветной XLSX.

CSV с весов — голые цифры, в них ничего не видно. Скрипт делает из него книгу
с подсветкой: вес цветовой шкалой, температура сине-красным градиентом, батарея
с полоской заряда и красной отсечкой на пороге алерта, привес день к дню
зелёным/красным. Плюс лист сводки, лист по дням и три графика.

Запуск:
    py -3 tools/csv_to_xlsx.py "C:\\path\\beehive_2026-08-13_2026-09-11.csv"
    py -3 tools/csv_to_xlsx.py          # возьмёт свежий beehive_*.csv с рабочего стола

Формат входа (как отдаёт прошивка): разделитель «;», десятичная запятая
    datetime;weight_kg;temp_c;humidity_pct;bat_v
"""
import sys
import csv
import glob
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─── Оформление ───────────────────────────────────────────────────────────
HDR_FILL = PatternFill("solid", fgColor="1F3A2E")
HDR_FONT = Font(bold=True, color="F5A623", size=11)
DAY_FILL = PatternFill("solid", fgColor="EFEFE7")   # полоса на первой строке новых суток
THIN = Side(style="thin", color="D0D0C8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def bat_percent(v):
    """Та же кусочно-линейная кривая LiPo, что в прошивке (Battery.cpp:53-59).
    Считаем по ней, иначе проценты в отчёте не сойдутся с показаниями весов."""
    if v >= 4.10:
        pct = 95 + (v - 4.10) / (4.20 - 4.10) * 5
    elif v >= 3.90:
        pct = 75 + (v - 3.90) / (4.10 - 3.90) * 20
    elif v >= 3.75:
        pct = 45 + (v - 3.75) / (3.90 - 3.75) * 30
    elif v >= 3.60:
        pct = 15 + (v - 3.60) / (3.75 - 3.60) * 30
    elif v >= 3.40:
        pct = 5 + (v - 3.40) / (3.60 - 3.40) * 10
    elif v >= 3.00:
        pct = (v - 3.00) / (3.40 - 3.00) * 5
    else:
        pct = 0
    return max(0, min(100, round(pct)))


def num(s):
    """«4,18» → 4.18 — прошивка пишет десятичную запятую под русский Excel."""
    s = (s or "").strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def read_rows(path):
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            raw = (r.get("datetime") or "").strip()
            dt = None
            for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M"):
                try:
                    dt = datetime.strptime(raw, fmt)
                    break
                except ValueError:
                    continue
            if dt is None:
                continue
            rows.append({
                "dt": dt,
                "w": num(r.get("weight_kg")),
                "t": num(r.get("temp_c")),
                "b": num(r.get("bat_v")),
            })
    rows.sort(key=lambda x: x["dt"])
    return rows


def style_header(ws, titles, widths):
    for i, (title, width) in enumerate(zip(titles, widths), start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def sheet_data(wb, rows):
    ws = wb.create_sheet("Замеры")
    style_header(
        ws,
        ["Дата и время", "Вес, кг", "Изм. к пред., кг", "Темп., °C", "Батарея, В", "Заряд, %"],
        [20, 11, 16, 11, 12, 10],
    )

    prev_w = None
    prev_day = None
    for i, r in enumerate(rows, start=2):
        new_day = prev_day is not None and r["dt"].date() != prev_day
        prev_day = r["dt"].date()

        ws.cell(row=i, column=1, value=r["dt"]).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(row=i, column=2, value=r["w"]).number_format = "0.00"

        delta = None if (prev_w is None or r["w"] is None) else round(r["w"] - prev_w, 2)
        dc = ws.cell(row=i, column=3, value=delta)
        dc.number_format = "+0.00;-0.00;0.00"
        if delta is not None and abs(delta) >= 0.05:
            dc.font = Font(bold=True, color="1E7B34" if delta > 0 else "C0392B")

        ws.cell(row=i, column=4, value=r["t"]).number_format = "0.0"
        ws.cell(row=i, column=5, value=r["b"]).number_format = "0.00"
        ws.cell(row=i, column=6, value=bat_percent(r["b"]) if r["b"] else None).number_format = '0"%"'

        for col in range(1, 7):
            c = ws.cell(row=i, column=col)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
            if new_day:
                c.fill = DAY_FILL

        if r["w"] is not None:
            prev_w = r["w"]

    last = len(rows) + 1
    # Вес — от светлого к насыщенному: сразу видно просадки и наборы
    ws.conditional_formatting.add(
        "B2:B%d" % last,
        ColorScaleRule(start_type="min", start_color="FFF3CD",
                       end_type="max", end_color="E8A33D"))
    # Температура — синий (холодно) → жёлтый → красный (жарко)
    ws.conditional_formatting.add(
        "D2:D%d" % last,
        ColorScaleRule(start_type="min", start_color="9DC3E6",
                       mid_type="percentile", mid_value=50, mid_color="FFF2CC",
                       end_type="max", end_color="E06666"))
    # Батарея — полоска заряда, и красным всё что ниже порога алерта 3.6 В
    ws.conditional_formatting.add(
        "E2:E%d" % last,
        DataBarRule(start_type="num", start_value=3.0, end_type="num",
                    end_value=4.2, color="63BE7B", showValue=True))
    ws.conditional_formatting.add(
        "E2:E%d" % last,
        CellIsRule(operator="lessThan", formula=["3.6"],
                   font=Font(bold=True, color="9C0006"),
                   fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.auto_filter.ref = "A1:F%d" % last
    return ws, last


def sheet_days(wb, rows):
    days = {}
    for r in rows:
        g = days.setdefault(r["dt"].date(), {"w": [], "t": [], "b": []})
        for k in ("w", "t", "b"):
            if r[k] is not None:
                g[k].append(r[k])

    ws = wb.create_sheet("По дням")
    style_header(
        ws,
        ["Дата", "Замеров", "Вес на конец, кг", "Привес за сутки, кг",
         "Вес мин – макс", "Темп. мин – макс, °C", "Батарея на конец, В"],
        [13, 10, 17, 19, 16, 20, 19],
    )

    prev_last = None
    for i, (d, g) in enumerate(sorted(days.items()), start=2):
        w_last = g["w"][-1] if g["w"] else None
        gain = None if (prev_last is None or w_last is None) else round(w_last - prev_last, 2)

        ws.cell(row=i, column=1, value=d).number_format = "DD.MM.YYYY"
        ws.cell(row=i, column=2, value=len(g["w"]))
        ws.cell(row=i, column=3, value=w_last).number_format = "0.00"
        gc = ws.cell(row=i, column=4, value=gain)
        gc.number_format = "+0.00;-0.00;0.00"
        if gain is not None and abs(gain) >= 0.05:
            gc.font = Font(bold=True, color="1E7B34" if gain > 0 else "C0392B")
        ws.cell(row=i, column=5,
                value="%.2f – %.2f" % (min(g["w"]), max(g["w"])) if g["w"] else "")
        ws.cell(row=i, column=6,
                value="%.1f – %.1f" % (min(g["t"]), max(g["t"])) if g["t"] else "")
        ws.cell(row=i, column=7, value=g["b"][-1] if g["b"] else None).number_format = "0.00"

        for col in range(1, 8):
            c = ws.cell(row=i, column=col)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")

        if w_last is not None:
            prev_last = w_last

    last = len(days) + 1
    ws.conditional_formatting.add(
        "D2:D%d" % last,
        ColorScaleRule(start_type="min", start_color="F8C9C4",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="max", end_color="A9D08E"))
    ws.auto_filter.ref = "A1:G%d" % last
    return ws, last


def add_charts(wb, ws_data, last_row):
    ws = wb.create_sheet("Графики")
    dates = Reference(ws_data, min_col=1, min_row=2, max_row=last_row)
    for title, col, color, anchor in (
        ("Вес улья, кг", 2, "E8A33D", "B2"),
        ("Температура, °C", 4, "C0504D", "B20"),
        ("Батарея, В", 5, "4F81BD", "B38"),
    ):
        ch = LineChart()
        ch.title = title
        ch.height, ch.width = 8, 26
        ch.y_axis.title = title
        ch.legend = None
        ch.add_data(Reference(ws_data, min_col=col, min_row=1, max_row=last_row),
                    titles_from_data=True)
        ch.set_categories(dates)
        ch.series[0].graphicalProperties.line.solidFill = color
        ch.series[0].graphicalProperties.line.width = 22000
        ch.series[0].smooth = False
        ws.add_chart(ch, anchor)
    return ws


def sheet_info(wb, rows, src):
    ws = wb.create_sheet("Сводка")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 46

    w = [r["w"] for r in rows if r["w"] is not None]
    t = [r["t"] for r in rows if r["t"] is not None]
    b = [r["b"] for r in rows if r["b"] is not None]
    span_days = max(1, (rows[-1]["dt"] - rows[0]["dt"]).days)

    # Расход считаем в процентах заряда, а не в вольтах: кривая LiPo нелинейна,
    # 0.1 В у верхушки и у дна — это совершенно разное количество энергии.
    used_pct = (bat_percent(b[0]) - bat_percent(b[-1])) if len(b) > 1 else 0
    per_day = used_pct / span_days if span_days else 0
    left_days = int(bat_percent(b[-1]) / per_day) if per_day > 0.01 else None
    full_cycle = int(100 / per_day) if per_day > 0.01 else None

    ws.cell(row=1, column=1, value="BeehiveScale — сводка по выгрузке").font = \
        Font(bold=True, size=14, color="1F3A2E")

    items = [
        ("Файл", os.path.basename(src)),
        ("Период", "%s — %s" % (rows[0]["dt"].strftime("%d.%m.%Y %H:%M"),
                                rows[-1]["dt"].strftime("%d.%m.%Y %H:%M"))),
        ("Дней наблюдения", span_days),
        ("Замеров", len(rows)),
        ("", ""),
        ("Вес: начало → конец", "%.2f → %.2f кг" % (w[0], w[-1])),
        ("Изменение за период", "%+.2f кг" % (w[-1] - w[0])),
        ("Вес мин / макс", "%.2f / %.2f кг" % (min(w), max(w))),
        ("", ""),
        ("Температура мин / макс", "%.1f / %.1f °C" % (min(t), max(t))),
        ("", ""),
        ("Батарея: начало → конец", "%.2f В (%d%%) → %.2f В (%d%%)"
            % (b[0], bat_percent(b[0]), b[-1], bat_percent(b[-1]))),
        ("Расход заряда", "%.0f%% за %d дн = %.2f%% в сутки" % (used_pct, span_days, per_day)),
        ("Осталось при таком темпе", ("≈ %d дней" % left_days) if left_days else "нет данных"),
        ("Полный цикл от 100%", ("≈ %d дней" % full_cycle) if full_cycle else "нет данных"),
    ]
    for i, (k, v) in enumerate(items, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color="3C4A32")
        ws.cell(row=i, column=2, value=v)
    return ws


def main():
    if len(sys.argv) > 1:
        src = sys.argv[1]
    else:
        found = sorted(
            glob.glob(os.path.join(os.path.expanduser("~"), "Desktop", "beehive_*.csv")),
            key=os.path.getmtime)
        if not found:
            print("CSV не найден. Укажи путь аргументом.")
            return 1
        src = found[-1]

    rows = read_rows(src)
    if not rows:
        print("В файле нет пригодных строк:", src)
        return 1

    wb = Workbook()
    wb.remove(wb.active)
    ws_data, last = sheet_data(wb, rows)
    sheet_days(wb, rows)
    add_charts(wb, ws_data, last)
    sheet_info(wb, rows, src)
    wb._sheets = [wb["Сводка"], wb["Замеры"], wb["По дням"], wb["Графики"]]

    dst = os.path.splitext(src)[0] + ".xlsx"
    wb.save(dst)
    print("Готово:", dst)
    print("Строк:", len(rows), "| дней:", len({r["dt"].date() for r in rows}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
